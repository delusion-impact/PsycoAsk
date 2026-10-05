"""Отчёт сравнения пресетов: тепловая карта, радар, бабочка, инсайты.

Режим «Сравнение» генерирует по анкете на каждый пресет и вместо отчёта
с data bars строит аналитический Excel (scripts/compare_presets.py).
"""

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

import compare_presets as report
import generate_data
import txt_to_qa

REPO = Path(__file__).resolve().parent.parent
MEM = REPO / "data" / "input" / "MEMpreset.json"
REF = REPO / "data" / "input" / "ref_txt.txt"


def _write_presets(path: Path, profiles: dict[str, list[int]]) -> Path:
    """profiles: имя пресета → 4 индекса в порядке ключей qa_fixture."""
    blocks = {name: {"демо": indexes[:2], "шкалы": indexes[2:]}
              for name, indexes in profiles.items()}
    path.write_text(json.dumps({"presets": blocks}, ensure_ascii=False),
                    encoding="utf-8")
    return path


# --- Типы вопросов, оценки, категории -------------------------------------

def test_classify_type():
    assert report.classify_type("Есть ли у вас дети?", ["Нет", "Да"]) == "Да-Нет"
    assert report.classify_type("Выберите утверждение, которое больше про вас:",
                                ["вариант А", "вариант Б"]) == "Выбор пары"
    assert report.classify_type("Мне трудно зависеть",
                                ["совсем не согласен", "Полностью согласен",
                                 "Затрудняюсь ответить"]) == "Шкала Лайкерта"
    assert report.classify_type("Ваш пол / Мужской", ["Мужской", None]) == "Демография"
    assert report.classify_type("Что-то ещё", ["вариант 1", "вариант 2"]) == "Прочее"


def test_score_likert_and_yesno():
    likert = ["совсем не согласен", "Скорее согласен",
              "Затрудняюсь ответить", "Полностью согласен"]
    assert report.score_answer("Шкала Лайкерта", likert, "Полностью согласен") == 1.0
    assert report.score_answer("Шкала Лайкерта", likert, "совсем не согласен") == 0.0
    assert report.score_answer("Шкала Лайкерта", likert, "Затрудняюсь ответить") == 0.5
    assert report.score_answer("Да-Нет", ["Да", "Нет", "Не знаю"], "Да") == 1.0
    assert report.score_answer("Да-Нет", ["Да", "Нет", "Не знаю"], "Нет") == 0.0
    assert report.score_answer("Да-Нет", ["Да", "Нет", "Не знаю"], "Не знаю") == 0.5
    assert report.score_answer("Демография", ["Мужской", None], "Мужской") is None
    assert report.score_answer("Шкала Лайкерта", likert, None) is None


def test_score_pair_is_option_position():
    options = ["первый вариант", "второй вариант"]
    assert report.score_answer("Выбор пары", options, options[0]) == 0.0
    assert report.score_answer("Выбор пары", options, options[1]) == 1.0


def test_band_colors():
    assert report.band(0.0) == "low"
    assert report.band(0.5) == "mid"
    assert report.band(1.0) == "high"
    assert report.band(None) is None


def test_category_rules_and_defaults():
    autonomy = "Мне трудно зависеть от других людей"
    assert report.category_of(autonomy, ["совсем не согласен", "Полностью согласен"],
                              "Шкала Лайкерта") == "Автономия vs Зависимость"
    assert report.category_of("Ваш пол / Мужской", ["Мужской", None],
                              "Демография") is None
    # Без ключевых слов — значение по умолчанию для типа вопроса
    assert report.category_of("Непонятный вопрос Да/Нет", ["Да", "Нет", "Не знаю"],
                              "Да-Нет") == "Коммуникация и секреты"


def test_block_labels():
    assert report.block_label("Шкала Лайкерта", 7, 181) == "Шкалы привязанности"
    assert report.block_label("Выбор пары", 40, 181) == "Бинарный выбор"
    assert report.block_label("Да-Нет", 106, 181) == "Да/Нет: партнёр"
    assert report.block_label("Да-Нет", 136, 181) == "Да/Нет: супруга"
    assert report.block_label("Да-Нет", 0, 4) == "Да/Нет"


# --- Сбор строк ------------------------------------------------------------

def test_build_rows_diff_and_demographics(qa_fixture):
    qa = json.loads(qa_fixture.read_text(encoding="utf-8"))
    # Европа / РФ / Китай: возраст совпадает у РФ и Китая, шкала расходится
    answers = [
        ["Мужской", None, "18-25", "Полностью согласен"],
        [None, "Женский", "26-35", "Скорее согласен"],
        [None, "Женский", "26-35", "Скорее не согласен"],
    ]
    profiles = [dict(zip(qa, item)) for item in answers]

    rows = report.build_rows(qa, profiles)
    by_type = {row["qtype"]: row for row in rows}
    likert = by_type["Шкала Лайкерта"]
    assert likert["scores"] == [1.0, 0.6, 0.25]
    assert likert["diff"] == pytest.approx(0.75)
    assert likert["block"] == "Шкалы привязанности"
    # Демография и «Прочее» не участвуют в разнице
    assert by_type["Демография"]["diff"] is None
    assert by_type["Прочее"]["diff"] is None


# --- Полный отчёт на синтетике --------------------------------------------

def test_report_builds_four_sheets(qa_fixture, tmp_path):
    preset_file = _write_presets(tmp_path / "mem.json", {
        "европейская семья 2026": [1, 1, 0, 0],
        "семья в рф 2026": [0, 1, 1, 1],
        "китайский вариант 2026": [1, 0, 1, 3],
    })
    out = report.build_report(output=tmp_path / "compare.xlsx",
                              qa_source=qa_fixture, preset_source=preset_file)
    assert out.exists()

    wb = load_workbook(out)
    assert wb.sheetnames == ["Тепловая карта", "Радар по категориям",
                             "Бабочка", "Инсайты топ-10"]

    heat = wb["Тепловая карта"]
    assert heat.max_row == 7  # заголовок + 4 вопроса
    assert heat.auto_filter.ref == "A3:H7"  # автофильтр = срезы openpyxl
    assert heat.freeze_panes == "A4"
    assert len(heat.conditional_formatting._cf_rules) == 1  # шкала «Разница»
    # Строки с ответами пресетов: эмодзи + текст
    likert_row = next(row for row in range(4, 8)
                      if heat.cell(row=row, column=6).value == "Шкала Лайкерта")
    assert str(heat.cell(row=likert_row, column=2).value).startswith("🟩")
    assert str(heat.cell(row=likert_row, column=4).value).startswith("🟥")

    radar = wb["Радар по категориям"]
    assert radar.max_row == 4  # одна категория + шапка
    assert radar.cell(row=4, column=1).value == "Эмоциональная близость"
    assert len(radar._charts) == 1

    insights = wb["Инсайты топ-10"]
    assert insights.max_row == 4  # только 1 вопрос с разницей
    assert insights.cell(row=4, column=6).value == pytest.approx(0.75)
    assert insights.cell(row=4, column=7).value == "⭐⭐"
    assert len(insights._charts) == 1


def test_report_rejects_wrong_index_count(qa_fixture, tmp_path):
    preset_file = _write_presets(tmp_path / "mem.json", {
        "европейская семья 2026": [1, 1, 0, 0],
        "семья в рф 2026": [0, 1, 1],       # не хватает индекса
        "китайский вариант 2026": [1, 0, 1, 3],
    })
    with pytest.raises(ValueError, match="индексов"):
        report.build_report(output=tmp_path / "compare.xlsx",
                            qa_source=qa_fixture, preset_source=preset_file)


# --- Три анкеты режима «Сравнение» ----------------------------------------

def test_comparison_records_one_per_preset(qa_fixture, tmp_path):
    indexes_by_preset = {
        "европейская семья 2026": [1, 1, 0, 0],
        "семья в рф 2026": [0, 1, 1, 1],
        "китайский вариант 2026": [1, 0, 1, 3],
    }
    preset_file = _write_presets(tmp_path / "mem.json", indexes_by_preset)
    records = generate_data.generate_comparison_records(
        qa_source=qa_fixture, preset_source=preset_file)

    assert len(records) == 3 == len(generate_data.COMPARISON_PRESETS)
    assert len({r["ID"] for r in records}) == 3
    qa = json.loads(qa_fixture.read_text(encoding="utf-8"))
    for name, record in zip(generate_data.COMPARISON_PRESETS, records):
        indexes = indexes_by_preset[name]
        expected = {key: options[i]
                    for (key, options), i in zip(qa.items(), indexes)}
        assert {k: v for k, v in record.items()
                if k not in ("ID", "Время создания")} == expected, name


# --- Реальные данные -------------------------------------------------------

@pytest.mark.skipif(not (MEM.exists() and REF.exists()),
                    reason="нет data/input/MEMpreset.json или ref_txt.txt")
def test_real_comparison_report(tmp_path):
    qa = txt_to_qa.build_qa(REF)
    qa_path = tmp_path / "qa.json"
    qa_path.write_text(json.dumps(qa, ensure_ascii=False), encoding="utf-8")

    out = report.build_report(output=tmp_path / "compare.xlsx",
                              qa_source=qa_path, preset_source=MEM)
    wb = load_workbook(out)

    heat = wb["Тепловая карта"]
    assert heat.max_row == len(qa) + 3
    assert heat.auto_filter.ref == f"A3:H{len(qa) + 3}"

    radar = wb["Радар по категориям"]
    categories = [radar.cell(row=row, column=1).value
                  for row in range(4, radar.max_row + 1)]
    assert len(categories) == 5
    assert set(categories) <= set(report.CATEGORIES)

    assert wb["Бабочка"]._charts
    insights = wb["Инсайты топ-10"]
    assert insights.max_row == 13  # шапка + 10 вопросов
    diffs = [insights.cell(row=row, column=6).value for row in range(4, 14)]
    assert diffs == sorted(diffs, reverse=True)
