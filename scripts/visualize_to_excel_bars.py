import json
from pathlib import Path

import pandas as pd
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from paths import BARS_XLSX, GENERATED_JSON, ensure_dirs, stamped

GENERATED_FILE = GENERATED_JSON
OUTPUT_XLSX = BARS_XLSX
EXCLUDE_KEYS = {"ID", "Время создания"}


def build_long_format(filepath):
    """
    Преобразует данные в 'длинный' формат:
    Вопрос | Вариант ответа | Процент
    Это необходимо для отрисовки отдельных баров для каждого варианта.
    """
    df = pd.read_json(filepath, orient='records')
    cols = [c for c in df.columns if c not in EXCLUDE_KEYS]

    rows = []
    for col in cols:
        # dropna=False — проценты считаем от всех записей вопроса, но пустые
        # ответы (None → NaN) в отчёт не попадают: без проверки pd.isna
        # подписью такого ответа становилась строка "nan"
        counts = df[col].value_counts(normalize=True, dropna=False)
        for answer, freq in counts.items():
            if pd.isna(answer):
                continue
            rows.append({
                "Вопрос": col,
                "Вариант ответа": str(answer),
                "Процент": round(freq * 100, 1)
            })

    return pd.DataFrame(rows, columns=["Вопрос", "Вариант ответа", "Процент"])


def style_and_format(writer, df):
    ws = writer.sheets["Распределение ответов"]

    # === 1. Стилизация заголовков ===
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

    # === 2. Ширина столбцов ===
    ws.column_dimensions['A'].width = 65  # Вопрос
    ws.column_dimensions['B'].width = 55  # Вариант ответа
    ws.column_dimensions['C'].width = 30  # Бар + процент

    # === 3. Выравнивание и границы для данных ===
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=3):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center', wrap_text=(cell.column == 1))
        # Центрируем текст процента поверх бара
        row[2].alignment = Alignment(horizontal='center', vertical='center')

    # === 4. Условное форматирование: Data Bars (Гистограммы) ===
    # Применяем к столбцу C (Процент)
    data_range = f"C2:C{ws.max_row}"

    rule = DataBarRule(
        start_type='num', start_value=0,
        end_type='num', end_value=100,
        color='4472C4',  # Синий бар (как в matplotlib палитре)
        showValue=True  # Показывать число поверх бара
    )
    ws.conditional_formatting.add(data_range, rule)

    # === 5. Группировка вопросов (визуальное разделение) ===
    # Добавляем легкую подложку для чередования вопросов
    current_q = None
    use_shade = False
    shade_fill = PatternFill(start_color="F2F7FB", end_color="F2F7FB", fill_type="solid")

    for row_idx in range(2, ws.max_row + 1):
        q_cell = ws.cell(row=row_idx, column=1)
        if q_cell.value != current_q:
            current_q = q_cell.value
            use_shade = not use_shade

        if use_shade:
            for col_idx in range(1, 4):
                ws.cell(row=row_idx, column=col_idx).fill = shade_fill

    # Фиксируем первую строку
    ws.freeze_panes = 'A2'


def build_bars_report(source=None, output=None) -> Path:
    """Таблица с Data Bars внутри ячеек: JSON после генерации → XLSX."""
    ensure_dirs()

    df_long = build_long_format(source or GENERATED_FILE)
    out_path = Path(output) if output else stamped(OUTPUT_XLSX)

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        df_long.to_excel(writer, sheet_name="Распределение ответов", index=False)
        style_and_format(writer, df_long)

    print(f"✅ Файл сохранён: {out_path}")
    print(f"   Всего строк: {len(df_long)}")
    print(f"   Визуализация: Data Bars (гистограммы внутри ячеек)")
    return out_path


if __name__ == "__main__":
    print("📊 Формирование таблицы с барами...")
    build_bars_report()